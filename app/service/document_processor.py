import asyncio
import hashlib
import logging
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import settings

logger = logging.getLogger(__name__)


@dataclass
class DocumentHierarchy:
    """Track document structure hierarchy."""
    level: int
    title: str
    parent: str | None = None
    page_number: int | None = None


@dataclass
class ChunkQuality:
    """Chunk quality metrics."""
    score: float
    length: int
    information_density: float
    is_boilerplate: bool
    reasons: list[str]


@dataclass
class PdfExtraction:
    """Result of a single-pass PDF extraction."""
    pages: list[tuple[int, str]]                       # (page_number, cleaned text)
    page_count: int
    hierarchy: list[DocumentHierarchy] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    images: list[dict[str, Any]] = field(default_factory=list)


class TextPreprocessor:
    """Text preprocessing and normalization for Persian and English."""

    BOILERPLATE_PATTERNS = [
        r"^\s*\d+\s*$",            # Page numbers only
        r"^صفحه\s*\d+\s*$",        # Persian page numbers
        r"^Page\s*\d+\s*$",
        r"^\s*[-_=]{3,}\s*$",      # Separator lines
        r"^\s*\.{3,}\s*$",         # Dots
        r"^[\s\d\-/]+$",           # Only numbers and separators
    ]

    ZWNJ = "\u200c"  # Zero-width non-joiner
    ZWJ = "\u200d"   # Zero-width joiner

    @classmethod
    def normalize_persian(cls, text: str) -> str:
        """Normalize Persian text Unicode issues."""
        text = text.replace("ي", "ی")  # Arabic yeh -> Persian yeh
        text = text.replace("ك", "ک")  # Arabic kaf -> Persian kaf
        text = re.sub(f"{cls.ZWNJ}+", cls.ZWNJ, text)
        text = re.sub(f"{cls.ZWJ}+", cls.ZWJ, text)
        return text

    @classmethod
    def normalize_whitespace(cls, text: str) -> str:
        """Normalize whitespace while preserving paragraph breaks."""
        text = text.replace("\t", " ")
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        lines = [line.rstrip() for line in text.split("\n")]
        text = "\n".join(lines)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" {2,}", " ", text)
        return text.strip()

    @classmethod
    def remove_headers_footers(cls, text: str) -> str:
        """Remove page-number style header/footer lines (single page)."""
        lines = text.split("\n")
        if len(lines) < 3:
            return text

        def looks_like_page_marker(line: str) -> bool:
            line = line.strip()
            return len(line) < 50 and (
                bool(re.match(r"^\d+$", line))
                or bool(re.search(r"صفحه|page", line, re.IGNORECASE))
            )

        if looks_like_page_marker(lines[0]):
            lines = lines[1:]
        if lines and looks_like_page_marker(lines[-1]):
            lines = lines[:-1]
        return "\n".join(lines)

    @classmethod
    def is_boilerplate(cls, text: str) -> bool:
        """Check if text is likely boilerplate."""
        text = text.strip()
        if len(text) < 10:
            return True
        return any(re.match(p, text, re.IGNORECASE) for p in cls.BOILERPLATE_PATTERNS)

    @classmethod
    def preprocess(cls, text: str) -> str:
        """Full preprocessing pipeline."""
        text = cls.normalize_whitespace(text)
        text = cls.normalize_persian(text)
        text = cls.remove_headers_footers(text)
        return text


class ChunkQualityScorer:
    """Score and filter document chunks."""

    MIN_CHUNK_LENGTH = 20
    MIN_QUALITY_SCORE = 0.3

    @classmethod
    def calculate_information_density(cls, text: str) -> float:
        """Ratio of alphanumeric characters to total characters."""
        if not text:
            return 0.0
        return sum(c.isalnum() for c in text) / len(text)

    @classmethod
    def score_chunk(cls, text: str) -> ChunkQuality:
        """Score a chunk's quality."""
        text = text.strip()
        length = len(text)
        reasons: list[str] = []

        is_boilerplate = TextPreprocessor.is_boilerplate(text)
        if is_boilerplate:
            reasons.append("boilerplate_pattern")

        if length < cls.MIN_CHUNK_LENGTH:
            reasons.append("too_short")

        density = cls.calculate_information_density(text)
        if density < 0.4:
            reasons.append("low_information_density")

        score = 1.0
        if is_boilerplate:
            score *= 0.1

        if length < cls.MIN_CHUNK_LENGTH:
            score *= 0.2
        elif length < 50:
            score *= 0.6

        if density < 0.3:
            score *= 0.3
        elif density < 0.5:
            score *= 0.7

        sentences = len(re.findall(r"[.!?؟]+", text))
        if sentences == 0 and length > 50:
            score *= 0.8
            reasons.append("no_sentence_endings")

        return ChunkQuality(
            score=score,
            length=length,
            information_density=density,
            is_boilerplate=is_boilerplate,
            reasons=reasons,
        )

    @classmethod
    def should_keep_chunk(cls, quality: ChunkQuality) -> bool:
        return quality.score >= cls.MIN_QUALITY_SCORE


class DocumentProcessor:
    """Advanced document processor for PDF, DOCX, and XLSX files."""

    # Heading heuristics (relative to the document's body font size, which is
    # measured as the char-weighted mode of span sizes; bold at body size is
    # NOT a heading — labels and table headers are routinely bold).
    MAX_HEADING_CHARS = 120
    HEADING_SIZE_DELTA = 0.8       # pt above body size -> heading candidate
    HEADING_LARGE_DELTA = 2.5      # pt above body size -> level-1 heading
    HEADING_BOLD_RATIO = 0.6       # min bold share for a modest size bump

    # Running header/footer removal
    EDGE_LINES = 2
    REPEATED_LINE_RATIO = 0.5
    MIN_PAGES_FOR_REPEAT_CHECK = 3

    # Prefer actual ruling lines to inferred text alignment. The latter needs
    # additional gutter checks so ordinary paragraphs do not become tables.
    TABLE_STRATEGIES: tuple[dict[str, Any], ...] = (
        {"vertical_strategy": "lines_strict", "horizontal_strategy": "lines_strict"},
        {
            "vertical_strategy": "lines_strict",
            "horizontal_strategy": "text",
            "snap_tolerance": 5,
        },
        {
            "vertical_strategy": "text",
            "horizontal_strategy": "text",
            "min_words_vertical": 3,
        },
    )

    def __init__(self) -> None:
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            separators=["\n\n", "\n", ".", "!", "?", "،", "؛", " ", ""],
            length_function=len,
        )
        self.upload_dir = Path("./uploads")
        self.upload_dir.mkdir(exist_ok=True)

    def __repr__(self) -> str:
        return (
            f"DocumentProcessor("
            f"chunk_size={settings.chunk_size}, "
            f"chunk_overlap={settings.chunk_overlap})"
        )

    # ------------------------------------------------------------------
    # ID generation
    # ------------------------------------------------------------------

    @staticmethod
    def generate_doc_id(file_content: bytes, file_name: str) -> str:
        """
        Generate a document ID from content hash + filename.
        Two uploads of the same bytes produce the same ID — intentional for
        deduplication. Callers must pass the original filename (not a
        randomized temp name), or dedup breaks.
        """
        return hashlib.sha256(file_content + file_name.encode()).hexdigest()[:24]

    # ------------------------------------------------------------------
    # Shared helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _rows_to_markdown(rows: list[list[Any]]) -> str:
        """Convert table rows (first row = header) to a markdown table."""

        def clean(cell: Any) -> str:
            if cell is None:
                return ""
            text = re.sub(r"\s+", " ", str(cell)).strip()
            return TextPreprocessor.normalize_persian(text).replace("|", "\\|")

        cleaned = [[clean(c) for c in row] for row in rows if row]
        cleaned = [row for row in cleaned if any(row)]
        if len(cleaned) < 2:
            return ""

        width = max(len(row) for row in cleaned)
        cleaned = [row + [""] * (width - len(row)) for row in cleaned]
        header, *body = cleaned

        lines = [
            "| " + " | ".join(header) + " |",
            "|" + "|".join(["---"] * width) + "|",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in body)
        return "\n".join(lines)

    # ------------------------------------------------------------------
    # Public entry point
    # ------------------------------------------------------------------

    async def process_file(
        self, file_path: str, original_name: Optional[str] = None
    ) -> tuple[list[Document], dict[str, Any]]:
        return await asyncio.to_thread(self._process_file, file_path, original_name)

    def _process_file(
        self, file_path: str, original_name: Optional[str] = None
    ) -> tuple[list[Document], dict[str, Any]]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        file_content = path.read_bytes()
        # Identity (doc_id, file_name metadata) follows the user's original
        # filename so re-uploads deduplicate; the temp path is irrelevant.
        file_name = original_name or path.name
        doc_id = self.generate_doc_id(file_content, file_name)
        ext = path.suffix.lower()

        if ext == ".pdf":
            return self._process_pdf(path, doc_id, file_name)
        if ext == ".docx":
            return self._process_docx(path, doc_id, file_name)
        if ext == ".xlsx":
            return self._process_xlsx(path, doc_id, file_name)
        raise ValueError(f"Unsupported file type: {ext}")

    # ------------------------------------------------------------------
    # PDF — single-pass extraction
    # ------------------------------------------------------------------

    def _extract_pdf_content(self, file_path: Path, doc_id: str) -> PdfExtraction:
        """
        Open the PDF once and, per page:
          1. detect valid tables (PyMuPDF for ruled tables, pdfplumber text
             alignment as fallback for borderless ones) and record bboxes,
          2. extract text lines that are NOT inside those tables,
          3. extract unique images (deduped by xref and content hash).
        Then remove running headers/footers across pages and classify
        headings relative to the document's body font size.
        """
        import pymupdf as fitz

        raw_pages: list[tuple[int, list[tuple[str, float, float]]]] = []
        hierarchy: list[DocumentHierarchy] = []
        tables: list[dict[str, Any]] = []
        images: list[dict[str, Any]] = []
        seen_xrefs: set[int] = set()
        seen_hashes: set[str] = set()
        images_dir = self.upload_dir / "images" / doc_id

        plumber_pdf = None

        def plumber_page(number: int):
            nonlocal plumber_pdf
            if plumber_pdf is None:
                import pdfplumber
                plumber_pdf = pdfplumber.open(str(file_path))
            return plumber_pdf.pages[number - 1]

        try:
            with fitz.open(str(file_path)) as doc:
                page_count = len(doc)

                toc = doc.get_toc()
                for level, title, page_num in toc:
                    title = TextPreprocessor.normalize_persian(" ".join(title.split()))
                    # PDF bookmarks can point to whole paragraphs or citations;
                    # their presence alone does not make them section headings.
                    if (level > 0 and 1 <= page_num <= page_count
                            and self._is_heading_title(title)):
                        hierarchy.append(DocumentHierarchy(
                            level=level, title=title, page_number=page_num,
                        ))
                # Fall back to font heuristics when no usable TOC entries remain.
                use_heuristic = not hierarchy

                for page_num, page in enumerate(doc, 1):
                    self._extract_page_images(
                        doc, page, page_num, images_dir, seen_xrefs, seen_hashes, images
                    )
                    table_rects = self._extract_page_tables(
                        page, page_num, tables, plumber_page
                    )
                    table_rects.extend(self._extract_captioned_image_tables(
                        page, page_num, images, tables, table_rects
                    ))
                    lines = self._extract_page_lines(page, table_rects)
                    raw_pages.append((page_num, lines))
        finally:
            if plumber_pdf is not None:
                plumber_pdf.close()

        raw_pages = self._remove_repeated_edge_lines(raw_pages)

        if use_heuristic:
            self._add_heuristic_headings(raw_pages, hierarchy)

        pages: list[tuple[int, str]] = []
        for page_num, lines in raw_pages:
            page_text = TextPreprocessor.preprocess(
                "\n".join(text for text, _, _ in lines)
            )
            if page_text.strip():
                pages.append((page_num, page_text))

        return PdfExtraction(
            pages=pages,
            page_count=page_count,
            hierarchy=hierarchy,
            tables=tables,
            images=images,
        )

    def _extract_page_tables(
        self,
        page: Any,
        page_num: int,
        out: list[dict[str, Any]],
        get_plumber_page: Callable[[int], Any],
    ) -> list[Any]:
        """Extract distinct physical tables; exclude only accepted regions."""
        import pymupdf as fitz

        rects: list[fitz.Rect] = []
        first_index = len(out)

        def accept_candidates(
            candidates: list[Any], words: list[dict[str, Any]] | None = None
        ) -> None:
            # Prefer a complete detection over its nested fragments regardless
            # of the order returned by the detector. Dedup is spatial and local
            # to this page: identical tables elsewhere are legitimate tables.
            regions = []
            for table in candidates:
                try:
                    rect = fitz.Rect(table.bbox)
                    if rect.is_empty or rect.is_infinite:
                        continue
                    regions.append((rect, table))
                except Exception as exc:
                    logger.debug("Invalid table region on page %d: %s", page_num, exc)
            regions.sort(key=lambda item: item[0].get_area(), reverse=True)
            for rect, table in regions:
                try:
                    if any(self._table_regions_overlap(rect, seen) for seen in rects):
                        continue
                    rows = table.extract()
                    if not self._is_valid_table(rows):
                        continue
                    if words is not None and not self._has_table_gutters(table, rows, words):
                        continue
                    markdown = self._rows_to_markdown(rows)
                    if not markdown:
                        continue
                except Exception as exc:
                    logger.debug(
                        "Failed to process table on page %d: %s", page_num, exc
                    )
                    continue
                rects.append(rect)
                out.append({
                    "content": markdown,
                    "page_number": page_num,
                    "table_index": len(out),
                })
            # Area determines deduplication preference; page position determines
            # output order. Keep table indices stable in reading order.
            ordered = sorted(zip(rects, out[first_index:]), key=lambda pair: (
                pair[0].y0, pair[0].x0,
            ))
            rects[:] = [rect for rect, _ in ordered]
            out[first_index:] = [record for _, record in ordered]
            for index, record in enumerate(out[first_index:], first_index):
                record["table_index"] = index

        # Strict lines ignore filled text boxes that can otherwise look like
        # ruled cells. Apply the same deduplication to both extraction engines.
        try:
            found = page.find_tables(strategy="lines_strict")
            accept_candidates(found.tables)
        except Exception as exc:
            logger.warning("Table detection failed on page %d: %s", page_num, exc)

        if rects:
            return rects

        try:
            plumber_page = get_plumber_page(page_num)
        except Exception as exc:
            logger.warning("pdfplumber fallback failed on page %d: %s", page_num, exc)
            return rects

        for table_settings in self.TABLE_STRATEGIES:
            try:
                candidates = plumber_page.find_tables(table_settings=table_settings)
                # Text alignment can split every word of a paragraph into a
                # cell. Require persistent whitespace between actual columns.
                words = (
                    plumber_page.extract_words()
                    if candidates and table_settings["vertical_strategy"] == "text"
                    else None
                )
                accept_candidates(candidates, words)
            except Exception as exc:
                logger.debug(
                    "Table strategy %r failed on page %d: %s",
                    table_settings, page_num, exc,
                )
                continue
            if rects:
                break  # first strategy with valid tables wins for this page
        return rects

    @staticmethod
    def _table_regions_overlap(first: Any, second: Any) -> bool:
        """Match near duplicates and nested fragments, including small ones."""
        intersection = (first & second).get_area()
        smaller = min(first.get_area(), second.get_area())
        union = first.get_area() + second.get_area() - intersection
        return smaller > 0 and (
            intersection / smaller >= 0.8 or intersection / union >= 0.5
        )

    @staticmethod
    def _has_table_gutters(
        table: Any, rows: list[list[Any]], words: list[dict[str, Any]]
    ) -> bool:
        """Require a repeated column gap wider than ordinary word spacing.

        Only text-inferred tables need this check. Empty spacer rows produced
        by pdfplumber are ignored, and the gap is measured in page coordinates
        so it works for both left-to-right and right-to-left text.
        """
        supported_gaps: Counter[tuple[int, int]] = Counter()
        observed_gaps: Counter[tuple[int, int]] = Counter()
        cell_lengths: dict[int, list[int]] = {}
        word_count = 0
        crossing_words = 0
        data_rows = 0
        for row, geometry in zip(rows, table.rows):
            occupied: list[tuple[int, list[dict[str, Any]]]] = []
            for column, (text, bbox) in enumerate(zip(row, geometry.cells)):
                if text is None or not str(text).strip() or bbox is None:
                    continue
                x0, top, x1, bottom = bbox
                cell_words = [
                    word for word in words
                    if x0 <= (word["x0"] + word["x1"]) / 2 < x1
                    and top <= (word["top"] + word["bottom"]) / 2 < bottom
                ]
                if cell_words:
                    occupied.append((column, cell_words))
                    cell_lengths.setdefault(column, []).append(len(cell_words))
                    word_count += len(cell_words)
                    crossing_words += sum(
                        word["x0"] < x0 - 1 or word["x1"] > x1 + 1
                        for word in cell_words
                    )
            if len(occupied) < 2:
                continue
            data_rows += 1
            for (left_col, left), (right_col, right) in zip(occupied, occupied[1:]):
                observed_gaps[left_col, right_col] += 1
                gap = min(word["x0"] for word in right) - max(word["x1"] for word in left)
                heights = [word["bottom"] - word["top"] for word in left + right]
                min_gap = max(6.0, sum(heights) / len(heights) * 0.6)
                if gap >= min_gap:
                    supported_gaps[left_col, right_col] += 1
        # False column boundaries often cut through words. A real gutter in a
        # two-column article is also insufficient: both sides contain prose.
        if word_count and crossing_words / word_count > 0.08:
            return False
        prose_columns = {
            column for column, lengths in cell_lengths.items()
            if sum(length >= 7 for length in lengths) > len(lengths) / 2
        }
        persistent = {
            pair for pair, count in observed_gaps.items()
            if count >= max(3, math.ceil(data_rows * 0.6))
        }
        genuine = {
            (left, right) for left, right in persistent
            if supported_gaps[left, right] >= math.ceil(observed_gaps[left, right] * 0.6)
            and not (left in prose_columns and right in prose_columns)
        }
        # One page gutter cannot validate dozens of inferred "columns" that
        # are just individual words. Most persistent boundaries must be real.
        return bool(persistent) and len(genuine) >= math.ceil(len(persistent) * 0.6)

    def _extract_captioned_image_tables(
        self,
        page: Any,
        page_num: int,
        images: list[dict[str, Any]],
        out: list[dict[str, Any]],
        text_table_rects: list[Any],
    ) -> list[Any]:
        """Preserve captioned raster tables without inventing their cell text.

        Match a standalone table label to a nearby image in the same column.
        Charts, portraits, and inline references to tables do not qualify.
        """
        import pymupdf as fitz

        text_lines = []
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                text = " ".join(span["text"].strip() for span in line["spans"]).strip()
                if text:
                    text_lines.append((fitz.Rect(line["bbox"]), text))
        captions = [
            (rect, text) for rect, text in text_lines
            if re.fullmatch(
                r"(?:table|جدول)\s+(?:[IVXLCDM]+|\d+)(?:\s*[:.\-–]\s*[^\n]{1,100})?",
                text, re.IGNORECASE,
            )
        ]
        if not captions:
            return []
        rects: list[Any] = []
        used_captions: set[int] = set()
        image_paths = {image["xref"]: image["image_path"] for image in images}
        for image in page.get_image_info(xrefs=True):
            image_rect = fitz.Rect(image["bbox"])
            if image_rect.is_empty or any(
                self._table_regions_overlap(image_rect, seen)
                for seen in text_table_rects + rects
            ):
                continue
            nearby = []
            for index, (caption_rect, caption) in enumerate(captions):
                if index in used_captions:
                    continue
                center_x = (caption_rect.x0 + caption_rect.x1) / 2
                if not image_rect.x0 <= center_x <= image_rect.x1:
                    continue
                distance = max(image_rect.y0 - caption_rect.y1,
                               caption_rect.y0 - image_rect.y1)
                if 0 <= distance <= 60:
                    nearby.append((distance, index, caption_rect, caption))
            if not nearby:
                continue
            _, index, caption_rect, caption = min(nearby, key=lambda match: match[0])
            used_captions.add(index)
            region = image_rect | caption_rect
            title_lines = []
            # Include a title between an above-table label and the image.
            if caption_rect.y1 <= image_rect.y0:
                for rect, text in sorted(text_lines, key=lambda item: item[0].y0):
                    if (rect.y0 >= caption_rect.y1 - 1 and rect.y1 <= image_rect.y0 + 1
                            and image_rect.x0 <= (rect.x0 + rect.x1) / 2 <= image_rect.x1):
                        title_lines.append(text)
                        region |= rect
            content = "\n".join([caption, *title_lines])
            content += "\n\n[Embedded table image; cell text requires OCR.]"
            out.append({
                "content": content,
                "page_number": page_num,
                "table_index": len(out),
                "extraction_method": "image",
                "requires_ocr": True,
                "image_path": image_paths.get(image.get("xref"), ""),
            })
            rects.append(region)
        return rects

    def _extract_page_lines(
        self, page: Any, table_rects: list[Any]
    ) -> list[tuple[str, float, float]]:
        """Return (text, max font size, bold ratio) for lines outside tables."""
        import pymupdf as fitz

        lines: list[tuple[str, float, float]] = []
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                if table_rects:
                    x0, y0, x1, y1 = line["bbox"]
                    center = fitz.Point((x0 + x1) / 2, (y0 + y1) / 2)
                    if any(rect.contains(center) for rect in table_rects):
                        continue

                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue

                texts = [s["text"].strip() for s in spans]
                total_chars = sum(len(t) for t in texts)
                if total_chars == 0:
                    continue

                text = " ".join(texts)
                size = max(float(s.get("size", 0) or 0) for s in spans)
                bold_chars = sum(
                    len(t) for s, t in zip(spans, texts) if s.get("flags", 0) & 16
                )
                lines.append((text, size, bold_chars / total_chars))
        return lines

    def _add_heuristic_headings(
        self,
        raw_pages: list[tuple[int, list[tuple[str, float, float]]]],
        hierarchy: list[DocumentHierarchy],
    ) -> None:
        """Classify headings relative to the body font size and append them.

        The body size is the char-weighted mode of all line sizes, so a
        document set in 11pt with a few 16pt titles yields ~11pt body and
        only the titles as headings. Repeated titles (running heads that
        survived edge removal) are deduplicated globally.
        """
        size_weights: dict[float, int] = {}
        for _, lines in raw_pages:
            for text, size, _ in lines:
                key = round(size * 2) / 2  # bucket to 0.5pt to ignore noise
                size_weights[key] = size_weights.get(key, 0) + len(text)
        if not size_weights:
            return
        body_size = max(size_weights, key=size_weights.get)

        seen_titles = {h.title for h in hierarchy}
        for page_num, lines in raw_pages:
            for text, size, bold_ratio in lines:
                level = self._heading_level(text, size, bold_ratio, body_size)
                if level is None:
                    continue
                title = TextPreprocessor.normalize_persian(text.strip())
                if not title or title in seen_titles:
                    continue
                seen_titles.add(title)
                hierarchy.append(DocumentHierarchy(
                    level=level, title=title, page_number=page_num,
                ))

    @classmethod
    def _heading_level(
        cls, text: str, size: float, bold_ratio: float, body_size: float
    ) -> Optional[int]:
        """Heading level (1 or 2) when a line stands out from body text."""
        if body_size <= 0 or size < body_size + cls.HEADING_SIZE_DELTA:
            return None
        large = size >= body_size + cls.HEADING_LARGE_DELTA
        if not large and bold_ratio < cls.HEADING_BOLD_RATIO:
            return None
        if not cls._is_heading_title(text):
            return None
        return 1 if large else 2

    @classmethod
    def _is_heading_title(cls, text: str) -> bool:
        """Common sanity checks for both bookmarks and inferred headings."""
        stripped = text.strip()
        if not (2 <= len(stripped) <= cls.MAX_HEADING_CHARS):
            return False
        if stripped[-1] in ".!؟?؛;،,":
            return False
        if re.fullmatch(r"[\d\s./:%-]+", stripped):
            return False
        if re.match(r"^\[\d+\]", stripped):
            return False
        return True

    def _extract_page_images(
        self,
        doc: Any,
        page: Any,
        page_num: int,
        images_dir: Path,
        seen_xrefs: set[int],
        seen_hashes: set[str],
        out: list[dict[str, Any]],
    ) -> None:
        """Extract unique images from a page (dedup by xref and content hash)."""
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)

            try:
                base = doc.extract_image(xref)
                if not base or not base.get("image"):
                    continue
                image_bytes: bytes = base["image"]

                digest = hashlib.sha256(image_bytes).hexdigest()
                if digest in seen_hashes:
                    continue
                seen_hashes.add(digest)

                try:
                    rects = page.get_image_rects(xref)
                    bbox = tuple(rects[0]) if rects else None
                except Exception:
                    bbox = None

                images_dir.mkdir(parents=True, exist_ok=True)
                image_index = len(out)
                image_path = images_dir / f"page_{page_num}_img_{image_index}.{base['ext']}"
                image_path.write_bytes(image_bytes)

                out.append({
                    "page_number": page_num,
                    "xref": xref,
                    "image_path": str(image_path),
                    "image_index": image_index,
                    "bbox": bbox,
                    "width": base.get("width"),
                    "height": base.get("height"),
                })
            except Exception as exc:
                logger.warning(
                    "Failed to extract image xref=%d on page %d: %s", xref, page_num, exc
                )

    @staticmethod
    def _line_signature(line: str) -> str:
        """Normalize a line for header/footer comparison (digits ignored)."""
        return re.sub(r"[\d۰-۹٠-٩]+", "#", line.strip().lower())

    def _remove_repeated_edge_lines(
        self, raw_pages: list[tuple[int, list[tuple[str, float, float]]]]
    ) -> list[tuple[int, list[tuple[str, float, float]]]]:
        """Drop running headers/footers that repeat at page edges."""
        if len(raw_pages) < self.MIN_PAGES_FOR_REPEAT_CHECK:
            return raw_pages

        k = self.EDGE_LINES
        counts: Counter[str] = Counter()
        for _, lines in raw_pages:
            edge = lines[:k] + lines[-k:]
            counts.update(self._line_signature(text) for text, _, _ in edge)

        threshold = max(2, math.ceil(len(raw_pages) * self.REPEATED_LINE_RATIO))
        repeated = {sig for sig, c in counts.items() if c >= threshold}
        if not repeated:
            return raw_pages

        logger.debug("Removing %d repeated header/footer patterns", len(repeated))

        cleaned: list[tuple[int, list[tuple[str, float, float]]]] = []
        for page_num, lines in raw_pages:
            n = len(lines)
            kept = [
                (text, size, bold_ratio)
                for i, (text, size, bold_ratio) in enumerate(lines)
                if not (
                    (i < k or i >= n - k)
                    and self._line_signature(text) in repeated
                )
            ]
            cleaned.append((page_num, kept))
        return cleaned

    def _is_valid_table(self, table: list[list[Any]] | None) -> bool:
        """Reject tables that are likely diagram boxes or callouts.

        Empty spacer rows are common in borderless extraction. Ignore them,
        but require at least two rows and columns containing actual data.
        """
        if not table or len(table) < 2:
            return False

        table = [row for row in table if row and any(
            cell is not None and str(cell).strip() for cell in row
        )]
        if len(table) < 2:
            return False
        col_count = max(len(row) for row in table)
        if col_count < 2:
            return False

        non_empty = sum(
            1 for row in table for cell in row if cell is not None and str(cell).strip()
        )
        occupied_columns = {
            column for row in table for column, cell in enumerate(row)
            if cell is not None and str(cell).strip()
        }
        if len(occupied_columns) < 2:
            return False

        if non_empty / (len(table) * col_count) < 0.3:
            return False

        # Header + single long row = likely a callout box
        if len(table) == 2:
            data_row = table[1]
            avg_cell_len = sum(len(str(c or "")) for c in data_row) / max(len(data_row), 1)
            if avg_cell_len > 80:
                return False

        return True

    def _build_hierarchy_path(
        self, hierarchies: list[DocumentHierarchy], current_page: int
    ) -> str:
        """Build level-aware hierarchical path for the given page."""
        relevant = [
            h for h in hierarchies
            if h.page_number is not None and h.page_number <= current_page
        ]
        if not relevant:
            return ""

        stack: dict[int, str] = {}
        for h in relevant:
            for lvl in [l for l in stack if l >= h.level]:
                del stack[lvl]
            stack[h.level] = h.title

        return " > ".join(stack[lvl] for lvl in sorted(stack))

    def _process_pdf(
        self, file_path: Path, doc_id: str, file_name: str
    ) -> tuple[list[Document], dict[str, Any]]:
        extraction = self._extract_pdf_content(file_path, doc_id)
        hierarchies = extraction.hierarchy

        documents: list[Document] = []
        filtered_count = 0
        text_chunks_kept = 0
        path_cache: dict[int, str] = {}

        def hierarchy_for(page_num: int | None) -> str:
            if not page_num or not hierarchies:
                return ""
            if page_num not in path_cache:
                path_cache[page_num] = self._build_hierarchy_path(hierarchies, page_num)
            return path_cache[page_num]

        # Text: split page-by-page so page_number is exact
        for page_num, page_text in extraction.pages:
            for chunk in self.text_splitter.split_text(page_text):
                quality = ChunkQualityScorer.score_chunk(chunk)
                if not ChunkQualityScorer.should_keep_chunk(quality):
                    filtered_count += 1
                    logger.debug(
                        "Filtered chunk on page %d (score=%.2f, reasons=%s)",
                        page_num, quality.score, quality.reasons,
                    )
                    continue

                documents.append(Document(
                    page_content=chunk,
                    metadata={
                        "doc_id": doc_id,
                        "file_name": file_name,
                        "chunk_index": len(documents),
                        "content_type": "text",
                        "source": file_name,
                        "page_number": page_num,
                        "hierarchy_path": hierarchy_for(page_num),
                        "quality_score": quality.score,
                        "information_density": quality.information_density,
                    },
                ))
                text_chunks_kept += 1

        # Tables
        for t in extraction.tables:
            table_metadata = {
                key: t[key] for key in ("extraction_method", "requires_ocr", "image_path")
                if key in t
            }
            documents.append(Document(
                page_content=t["content"],
                metadata={
                    "doc_id": doc_id,
                    "file_name": file_name,
                    "page_number": t["page_number"],
                    "table_index": t["table_index"],
                    "content_type": "table",
                    "source": file_name,
                    "hierarchy_path": hierarchy_for(t["page_number"]),
                    **table_metadata,
                },
            ))

        # Uncaptioned images are kept on disk. Captioned image tables have a
        # searchable caption chunk above, explicitly marked as requiring OCR.

        if not extraction.tables and extraction.page_count:
            logger.info(
                "No tables detected in %s (pages=%d); tables may be scanned "
                "images or lack detectable structure",
                file_name, extraction.page_count,
            )

        stats = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "pdf",
            "page_count": extraction.page_count,
            "chunk_count": len(documents),
            "text_chunks": text_chunks_kept,
            "table_count": len(extraction.tables),
            "image_count": len(extraction.images),
            "filtered_chunks": filtered_count,
            "hierarchy_levels": max((h.level for h in hierarchies), default=0),
            "heading_count": len(hierarchies),
            "tables_requiring_ocr": sum(bool(t.get("requires_ocr")) for t in extraction.tables),
        }
        logger.info(
            "Processed PDF %s: %d chunks (%d tables, %d filtered, %d headings)",
            file_name, len(documents), len(extraction.tables),
            filtered_count, len(hierarchies),
        )
        return documents, stats

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    def _process_docx(
        self, file_path: Path, doc_id: str, file_name: str
    ) -> tuple[list[Document], dict[str, Any]]:
        import docx

        document = docx.Document(str(file_path))
        documents: list[Document] = []
        filtered_count = 0
        text_chunks_kept = 0
        table_chunks_kept = 0

        # Heading paragraphs -> full hierarchical path (ancestors + own title)
        heading_paths: dict[int, str] = {}
        heading_levels: list[int] = []
        for i, para in enumerate(document.paragraphs):
            style_name = para.style.name if para.style and para.style.name else ""
            if not style_name.startswith("Heading"):
                continue
            text = para.text.strip()
            if not text:
                continue
            level = int(style_name.split()[-1]) if style_name[-1].isdigit() else 1
            path = self._docx_heading_path(document.paragraphs, i, level)
            path.append(text)
            heading_paths[i] = " > ".join(path)
            heading_levels.append(level)

        # Collect paragraphs and tables in document order
        # python-docx exposes tables separately, so merge by walking the body XML
        from docx.oxml.ns import qn

        body = document.element.body
        para_idx = 0
        table_idx = 0
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                para = (
                    document.paragraphs[para_idx]
                    if para_idx < len(document.paragraphs)
                    else None
                )
                para_idx += 1
                if para is None:
                    continue
                text = para.text.strip()
                if not text:
                    continue
                text = TextPreprocessor.preprocess(text)
                if not text:
                    continue

                hierarchy_path = heading_paths.get(para_idx - 1, "")

                for chunk in self.text_splitter.split_text(text):
                    quality = ChunkQualityScorer.score_chunk(chunk)
                    if not ChunkQualityScorer.should_keep_chunk(quality):
                        filtered_count += 1
                        continue

                    documents.append(Document(
                        page_content=chunk,
                        metadata={
                            "doc_id": doc_id,
                            "file_name": file_name,
                            "chunk_index": len(documents),
                            "content_type": "text",
                            "source": file_name,
                            "hierarchy_path": hierarchy_path,
                            "quality_score": quality.score,
                            "information_density": quality.information_density,
                        },
                    ))
                    text_chunks_kept += 1

            elif child.tag == qn("w:tbl"):
                table = (
                    document.tables[table_idx]
                    if table_idx < len(document.tables)
                    else None
                )
                table_idx += 1
                if table is None:
                    continue
                rows = self._docx_table_rows(table)
                markdown = self._rows_to_markdown(rows)
                if not markdown:
                    continue
                documents.append(Document(
                    page_content=markdown,
                    metadata={
                        "doc_id": doc_id,
                        "file_name": file_name,
                        "table_index": table_idx - 1,
                        "content_type": "table",
                        "source": file_name,
                    },
                ))
                table_chunks_kept += 1

        stats = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "docx",
            "page_count": 0,
            "chunk_count": len(documents),
            "text_chunks": text_chunks_kept,
            "table_count": table_chunks_kept,
            "image_count": 0,
            "filtered_chunks": filtered_count,
            "hierarchy_levels": max(heading_levels, default=0),
            "heading_count": len(heading_paths),
        }
        logger.info(
            "Processed DOCX %s: %d chunks (%d tables, %d filtered)",
            file_name, len(documents), table_chunks_kept, filtered_count,
        )
        return documents, stats

    @staticmethod
    def _docx_heading_path(
        paragraphs: list, current_idx: int, current_level: int
    ) -> list[str]:
        """Walk backwards from current_idx to collect ancestor headings."""
        path: list[tuple[int, str]] = []
        target_level = current_level - 1
        for i in range(current_idx - 1, -1, -1):
            para = paragraphs[i]
            if not para.style or not para.style.name or not para.style.name.startswith("Heading"):
                continue
            text = para.text.strip()
            if not text:
                continue
            lvl_str = para.style.name.split()[-1]
            lvl = int(lvl_str) if lvl_str.isdigit() else 1
            if lvl <= target_level:
                path.append((lvl, text))
                target_level = lvl - 1
            if target_level == 0:
                break
        path.reverse()
        return [t for _, t in path]

    @staticmethod
    def _docx_table_rows(table: Any) -> list[list[str]]:
        """Extract table rows, handling merged cells without duplicating text."""
        from docx.oxml.ns import qn

        rows: list[list[str]] = []
        for row in table.rows:
            cells: list[str] = []
            seen_texts: set[str] = set()
            for cell in row.cells:
                text = TextPreprocessor.normalize_whitespace(cell.text.strip())
                # Merged cells repeat the same _tc element; deduplicate by content
                if text and text not in seen_texts:
                    cells.append(text)
                    seen_texts.add(text)
                elif not text:
                    cells.append("")
            if any(cells):
                rows.append(cells)
        return rows

    # ------------------------------------------------------------------
    # XLSX
    # ------------------------------------------------------------------

    def _process_xlsx(
        self, file_path: Path, doc_id: str, file_name: str
    ) -> tuple[list[Document], dict[str, Any]]:
        from openpyxl import load_workbook

        wb = load_workbook(str(file_path), read_only=True, data_only=True)
        documents: list[Document] = []
        sheet_count = 0

        try:
            for sheet in wb.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]

                if len(rows) < 2:
                    continue

                # Skip sheets that are mostly empty
                non_empty = sum(
                    1 for r in rows for c in r if c is not None and str(c).strip()
                )
                if non_empty < 3:
                    continue

                markdown = self._rows_to_markdown(rows)
                if not markdown:
                    continue

                content = f"جدول از برگه: {sheet.title}\n\n{markdown}"
                # Large sheets must be split, otherwise they exceed the
                # embedding model's input limit and the upload fails.
                pieces = (
                    self.text_splitter.split_text(content)
                    if len(content) > settings.chunk_size
                    else [content]
                )

                sheet_count += 1
                for piece_index, piece in enumerate(pieces):
                    documents.append(Document(
                        page_content=piece,
                        metadata={
                            "doc_id": doc_id,
                            "file_name": file_name,
                            "sheet_name": sheet.title,
                            "sheet_index": sheet_count - 1,
                            "content_type": "table",
                            "source": file_name,
                            "chunk_index": piece_index,
                            "total_chunks": len(pieces),
                            "row_count": len(rows),
                        },
                    ))
        finally:
            wb.close()

        stats = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "xlsx",
            "page_count": 0,
            "chunk_count": len(documents),
            "text_chunks": 0,
            "table_count": sheet_count,
            "image_count": 0,
            "filtered_chunks": 0,
            "sheet_count": sheet_count,
        }
        logger.info(
            "Processed XLSX %s: %d chunks from %d sheets",
            file_name, len(documents), sheet_count,
        )
        return documents, stats
