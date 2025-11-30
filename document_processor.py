import hashlib
import logging
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
import pandas as pd
import numpy as np
from datetime import datetime

# PDF processing
import PyPDF2
import pdfplumber
import fitz  # PyMuPDF

# Word processing
from docx import Document as DocxDocument

# Excel processing
import openpyxl

# LangChain
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.schema import Document
from config import settings

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Document processor for PDF, DOCX, and XLSX files"""
    
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            separators=["\n\n", "\n", ".", "!", "?", "،", "؛", " ", ""],
            length_function=len,
        )
        self.upload_dir = Path("./uploads")
        self.upload_dir.mkdir(exist_ok=True)
    
    def generate_doc_id(self, file_content: bytes, file_name: str) -> str:
        """Generate unique document ID"""
        content_hash = hashlib.sha256(file_content).hexdigest()[:16]
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        return f"{timestamp}_{content_hash}"
    
    def process_file(self, file_path: str) -> Tuple[List[Document], Dict[str, Any]]:
        """
        Process a file from path
        
        Args:
            file_path: Path to the file
            
        Returns:
            Tuple of (chunks, metadata)
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # Read file content
        with open(file_path, "rb") as f:
            file_content = f.read()
        
        file_name = file_path.name
        doc_id = self.generate_doc_id(file_content, file_name)
        
        # Determine file type and process
        ext = file_path.suffix.lower()
        
        if ext == ".pdf":
            return self._process_pdf(file_path, doc_id, file_name)
        elif ext == ".docx":
            return self._process_docx(file_path, doc_id, file_name)
        elif ext == ".xlsx":
            return self._process_xlsx(file_path, doc_id, file_name)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    
    def _extract_pdf_text(self, file_path: Path) -> Tuple[str, int]:
        """Extract text from PDF using PyMuPDF"""
        doc = fitz.open(str(file_path))
        text_parts = []
        page_count = len(doc)  # Get page count before closing
        
        for page_num, page in enumerate(doc, 1):
            page_text = page.get_text()
            if page_text:
                text_parts.append(f"[صفحه {page_num}]\n{page_text}")
        
        doc.close()
        return "\n".join(text_parts), page_count  # Return the saved page_count
    
    def _extract_pdf_tables(self, file_path: Path) -> List[Dict[str, Any]]:
        """Extract tables from PDF using pdfplumber"""
        tables_data = []
        
        try:
            with pdfplumber.open(str(file_path)) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    tables = page.extract_tables()
                    for table_idx, table in enumerate(tables):
                        if table and len(table) > 1:  # Ensure table has data
                            try:
                                # Clean table data
                                cleaned_table = []
                                for row in table:
                                    if row:  # Skip empty rows
                                        cleaned_row = [str(cell) if cell else "" for cell in row]
                                        cleaned_table.append(cleaned_row)
                                
                                if len(cleaned_table) > 1:  # Need header + at least one row
                                    df = pd.DataFrame(cleaned_table[1:], columns=cleaned_table[0])
                                    markdown_table = df.to_markdown(index=False)
                                    
                                    tables_data.append({
                                        "content": markdown_table,
                                        "page_number": page_num,
                                        "table_index": table_idx
                                    })
                            except Exception as e:
                                logger.warning(f"Failed to process table {table_idx} on page {page_num}: {e}")
        except Exception as e:
            logger.warning(f"Failed to extract tables from PDF: {e}")
        
        return tables_data
    
    def _process_pdf(self, file_path: Path, doc_id: str, file_name: str) -> Tuple[List[Document], Dict[str, Any]]:
        """Process PDF file"""
        # Extract text
        text_content, page_count = self._extract_pdf_text(file_path)
        
        # Extract tables
        tables = self._extract_pdf_tables(file_path)
        
        # Create chunks from text
        text_chunks = []
        if text_content.strip():  # Only split if there's actual text
            text_chunks = self.text_splitter.split_text(text_content)
        
        documents = []
        
        # Add text chunks
        for i, chunk in enumerate(text_chunks):
            documents.append(Document(
                page_content=chunk,
                metadata={
                    "doc_id": doc_id,
                    "file_name": file_name,
                    "chunk_index": i,
                    "content_type": "text",
                    "source": str(file_path)
                }
            ))
        
        # Add table chunks
        for table_data in tables:
            documents.append(Document(
                page_content=table_data["content"],
                metadata={
                    "doc_id": doc_id,
                    "file_name": file_name,
                    "page_number": table_data["page_number"],
                    "content_type": "table",
                    "source": str(file_path)
                }
            ))
        
        metadata = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "pdf",
            "page_count": page_count,
            "chunk_count": len(documents),
            "table_count": len(tables),
            "text_chunks": len(text_chunks)
        }
        
        logger.info(f"Processed PDF: {file_name} - {len(documents)} total chunks")
        return documents, metadata
    
    def _process_docx(self, file_path: Path, doc_id: str, file_name: str) -> Tuple[List[Document], Dict[str, Any]]:
        """Process Word DOCX file"""
        doc = DocxDocument(str(file_path))
        
        text_parts = []
        tables_data = []
        
        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        
        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            try:
                rows = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    # Check if row has any data
                    if any(row_data):
                        rows.append(row_data)
                
                if len(rows) > 1:  # Need header + at least one row
                    df = pd.DataFrame(rows[1:], columns=rows[0])
                    markdown_table = df.to_markdown(index=False)
                    tables_data.append(markdown_table)
            except Exception as e:
                logger.warning(f"Failed to process table {table_idx}: {e}")
        
        # Combine and chunk text
        full_text = "\n".join(text_parts)
        text_chunks = []
        if full_text.strip():  # Only split if there's text
            text_chunks = self.text_splitter.split_text(full_text)
        
        documents = []
        
        # Add text chunks
        for i, chunk in enumerate(text_chunks):
            documents.append(Document(
                page_content=chunk,
                metadata={
                    "doc_id": doc_id,
                    "file_name": file_name,
                    "chunk_index": i,
                    "content_type": "text",
                    "source": str(file_path)
                }
            ))
        
        # Add table chunks
        for i, table_md in enumerate(tables_data):
            documents.append(Document(
                page_content=table_md,
                metadata={
                    "doc_id": doc_id,
                    "file_name": file_name,
                    "table_index": i,
                    "content_type": "table",
                    "source": str(file_path)
                }
            ))
        
        metadata = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "docx",
            "chunk_count": len(documents),
            "table_count": len(tables_data),
            "text_chunks": len(text_chunks)
        }
        
        logger.info(f"Processed DOCX: {file_name} - {len(documents)} total chunks")
        return documents, metadata
    
    def _process_xlsx(self, file_path: Path, doc_id: str, file_name: str) -> Tuple[List[Document], Dict[str, Any]]:
        """Process Excel XLSX file"""
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        
        documents = []
        sheet_count = len(wb.sheetnames)
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            
            # Convert sheet to DataFrame
            data = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out empty rows
                if any(cell is not None for cell in row):
                    data.append([str(cell) if cell is not None else "" for cell in row])
            
            if len(data) > 1:  # Need header + at least one data row
                try:
                    df = pd.DataFrame(data[1:], columns=data[0])
                    markdown_table = df.to_markdown(index=False)
                    
                    documents.append(Document(
                        page_content=f"جدول از برگه: {sheet_name}\n\n{markdown_table}",
                        metadata={
                            "doc_id": doc_id,
                            "file_name": file_name,
                            "sheet_name": sheet_name,
                            "content_type": "table",
                            "source": str(file_path)
                        }
                    ))
                except Exception as e:
                    logger.warning(f"Failed to process sheet {sheet_name}: {e}")
        
        wb.close()
        
        metadata = {
            "doc_id": doc_id,
            "file_name": file_name,
            "doc_type": "xlsx",
            "sheet_count": sheet_count,
            "chunk_count": len(documents),
            "table_count": len(documents)
        }
        
        logger.info(f"Processed XLSX: {file_name} - {len(documents)} total chunks")
        return documents, metadata